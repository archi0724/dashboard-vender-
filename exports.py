"""Formatted, formula-safe Excel and CSV exports."""
from __future__ import annotations
from io import BytesIO
import math
import re

import pandas as pd

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.comments import Comment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    OPENPYXL_IMPORT_ERROR = None
except ModuleNotFoundError as error:
    Workbook = None
    OPENPYXL_IMPORT_ERROR = error

from vendor_core import DOCUMENT_TYPES, supporting_category


def csv_bytes(frame: pd.DataFrame) -> bytes:
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + value
        return value
    return frame.map(safe).to_csv(index=False).encode("utf-8-sig")


def write_text(cell, value):
    cell.value = "" if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))
    cell.data_type = "s"  # A vendor/file name must never become an Excel formula.


def style_sheet(sheet, headers: list[str], widths: list[int], title: str, note: str):
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    write_text(sheet.cell(1, 1), title)
    sheet.cell(1, 1).font = Font(name="Calibri", size=18, bold=True, color="17221B")
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(2, 1, note).font = Font(name="Calibri", size=10, color="667268")
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34
    for i, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(4, i, header)
        cell.fill = PatternFill("solid", fgColor="233B2E")
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.row_dimensions[4].height = 32
    sheet.freeze_panes = "B5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"


def finish(sheet):
    sheet.auto_filter.ref = f"A4:{get_column_letter(sheet.max_column)}{max(sheet.max_row,4)}"
    for row in sheet.iter_rows(min_row=5):
        height = 24
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="000000" if cell.data_type == "f" else "24613C")
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left" if cell.column == 1 else "center")
            if cell.value == "Yes":
                cell.fill = PatternFill("solid", fgColor="E3F0D8")
            elif cell.value == "No":
                cell.fill = PatternFill("solid", fgColor="FBE9E7")
            elif cell.row % 2:
                cell.fill = PatternFill("solid", fgColor="F5F7F2")
            if isinstance(cell.value, str) and cell.data_type != "f":
                width = sheet.column_dimensions[cell.column_letter].width or 15
                height = max(height, min(90, 15 * (math.ceil(len(cell.value) / max(width - 2, 1)) + 1)))
        sheet.row_dimensions[row[0].row].height = height


def workbook_bytes(checklist: pd.DataFrame, documents: pd.DataFrame, per_company: bool = False, aliases: pd.DataFrame | None = None) -> bytes:
    if OPENPYXL_IMPORT_ERROR is not None:
        raise RuntimeError("Excel export requires the openpyxl package. Add openpyxl to requirements.txt and redeploy the app.") from OPENPYXL_IMPORT_ERROR
    book = Workbook()
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    sheet = book.active
    sheet.title = "Document Checklist"
    headers = ["Company Name", *DOCUMENT_TYPES, "Available", "Missing", "Completion", "Files", "Needs review"]
    available_column = 2 + len(DOCUMENT_TYPES)
    missing_column = available_column + 1
    completion_column = missing_column + 1
    note = "Yes = an available file is classified to this category. No = no classified file found. This is NOT validity/compliance verification."
    style_sheet(sheet, headers, [44] + [17] * len(DOCUMENT_TYPES) + [13, 13, 15, 12, 16], (str(checklist.iloc[0]["Company Name"]) + " - DOCUMENT CHECKLIST") if len(checklist)==1 else "VENDOR DOCUMENT CHECKLIST", note)
    for r, record in enumerate(checklist.to_dict("records"), 5):
        for c, header in enumerate(headers, 1):
            value = record[header]
            if header in ("Company Name", *DOCUMENT_TYPES):
                write_text(sheet.cell(r,c), value)
            else:
                sheet.cell(r,c, value)
        first_type = get_column_letter(2)
        last_type = get_column_letter(1 + len(DOCUMENT_TYPES))
        sheet.cell(r, available_column, f'=COUNTIF({first_type}{r}:{last_type}{r},"Yes")')
        sheet.cell(r, missing_column, f'=COLUMNS({first_type}{r}:{last_type}{r})-{get_column_letter(available_column)}{r}')
        sheet.cell(r, completion_column, f'=IFERROR({get_column_letter(available_column)}{r}/COLUMNS({first_type}{r}:{last_type}{r}),0)').number_format = "0%"
        sheet.cell(r,1).comment = Comment("Source: company folder or uploaded vendor master. See Document Register for source paths and classification evidence.", "Vendor Document Desk")
    finish(sheet)
    register = book.create_sheet("Document Register")
    register_headers = ["Company Name", "Document Type", "Filename", "File available", "Review", "Classification", "Source path", "Upload batch", "Uploaded at", "Reason"]
    style_sheet(register, register_headers, [40, 25, 72, 16, 16, 25, 90, 42, 28, 82], "DOCUMENT REGISTER", "Original source paths are retained for audit only. Handover names are never appended to Company Name.")
    for r, doc in enumerate(documents.itertuples(index=False), 5):
        values = [doc.company_name or "Unassigned", ", ".join(doc.types) or supporting_category(doc.filename) or "Other / Needs review", doc.filename,
                  "Yes" if doc.available else "No", "Needs review" if doc.needs_review or not doc.available else "Supporting document" if not doc.types else "Classified",
                  doc.method, doc.original_path, doc.source_batch, doc.uploaded_at, doc.reason]
        for c, value in enumerate(values,1):
            write_text(register.cell(r,c), value)
    finish(register)
    master = book.create_sheet("Company Master", 0)
    master_headers = ["Canonical ID", "Canonical Company Name", "Match Status"]
    style_sheet(master, master_headers, [18, 52, 24], "COMPANY MASTER", "One canonical row per company. IDs remain stable for future uploads.")
    for r, record in enumerate(checklist.to_dict("records"), 5):
        values = [record.get("canonical_id", ""), record["Company Name"], "Confirmed"]
        for c, value in enumerate(values, 1):
            write_text(master.cell(r, c), value)
    finish(master)
    alias_sheet = book.create_sheet("Alias Mapping", 1)
    alias_headers = ["Canonical ID", "Canonical Company", "Original / Detected Name", "Source", "Confidence", "Status"]
    style_sheet(alias_sheet, alias_headers, [18, 42, 42, 42, 14, 24], "COMPANY ALIAS MAPPING", "Aliases are retained for future matching. Uncertain matches are not silently merged.")
    alias_rows = aliases.to_dict("records") if aliases is not None and not aliases.empty else []
    canonical_ids = {record["company_key"]: record.get("canonical_id", "") for record in checklist.to_dict("records")}
    canonical_names = {record["company_key"]: record["Company Name"] for record in checklist.to_dict("records")}
    for r, alias in enumerate(alias_rows, 5):
        values = [canonical_ids.get(alias["canonical_key"], ""), canonical_names.get(alias["canonical_key"], ""),
                  alias["alias_name"], alias["source"], f'{int(alias["confidence"])}%', alias["status"]]
        for c, value in enumerate(values, 1):
            write_text(alias_sheet.cell(r, c), value)
    finish(alias_sheet)
    if per_company:
        names_used = {name.casefold() for name in book.sheetnames}
        for record in checklist.to_dict("records"):
            company = record["Company Name"]
            base = re.sub(r"[\\/*?:\[\]]", " ", company).strip(" '")[:31] or "Company"
            name, suffix = base, 2
            while name.casefold() in names_used:
                tail = f" ({suffix})"; name = base[:31-len(tail)] + tail; suffix += 1
            names_used.add(name.casefold())
            tab = book.create_sheet(name)
            style_sheet(tab, ["Document Type", "Available", "Stored files"], [26, 16, 95], company, note)
            group = documents[documents.company_key == record["company_key"]]
            for r, category in enumerate(DOCUMENT_TYPES,5):
                names = [doc.filename for doc in group.itertuples(index=False) if category in doc.types and doc.available]
                for c, value in enumerate([category, record[category], "\n".join(names) or "No classified file available"],1):
                    write_text(tab.cell(r,c), value)
            other = group[group.types.map(lambda value: not value)]
            for r, doc in enumerate(other.itertuples(index=False), 5+len(DOCUMENT_TYPES)):
                values = [supporting_category(doc.filename) or "Other / Needs review", "Yes" if doc.available else "No", doc.filename]
                for c, value in enumerate(values, 1):
                    write_text(tab.cell(r,c), value)
            finish(tab)
    readme = book.create_sheet("Read me")
    notes = [
        ["Meaning", "Yes means file presence, not authenticity, expiry, completeness or legal compliance."],
        ["No", "No matching available file was found. Check Needs review before concluding the document was never supplied."],
        ["Company ownership", "Company folders take precedence over names inside file names. Generic handover folders are ignored."],
        ["Udyam", "Includes Udyam, Udhyam, MSME and Udyog Aadhaar business registration; not personal Aadhaar."],
        ["ASF ISO", "All recognized ISO certificate names map to the original ASF ISO column, regardless of vendor name."],
        ["MD", "MD-form/medical-device licence labels; long certificate serials beginning MD are not MD licences."],
        ["Other files", "Recognised supporting documents appear in the register and company sheet. Only unclear/unavailable files need review. All files are retained."],
        ["Head count", "One row per unique company. Files counts stored document records, not document types. Repeated identical file bytes within one company count once."],
        ["Privacy", "This export may contain sensitive file names. Share only with authorized people."],
        ["Canonical IDs", "Company Master assigns one stable COMP-### ID per company. Alias Mapping preserves detected names and source context."],
        ["Uncertain matches", "Possible matches remain separate and should be corrected through Review files after identifying evidence is supplied."],
    ]
    style_sheet(readme,["Item","Definition"],[26,115],"HOW TO READ THIS WORKBOOK","Source: uploaded vendor master and uploaded document files; no external data or third-party classification API.")
    for r, values in enumerate(notes,5):
        for c, value in enumerate(values,1):write_text(readme.cell(r,c),value)
    finish(readme)
    output = BytesIO(); book.save(output)
    return output.getvalue()
