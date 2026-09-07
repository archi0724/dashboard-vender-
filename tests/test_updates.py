from io import BytesIO
import hashlib
import json
import sqlite3
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook
from import_service import import_documents
from vendor_core import classify, dashboard_counts, export_filename, build_checklist
from storage import Store
from exports import workbook_bytes


class Upload:
    def __init__(self,entries):
        self.name='Person vendor.zip'; out=BytesIO()
        with zipfile.ZipFile(out,'w') as z:
            for name,data in entries:z.writestr(name,data)
        self.data=out.getvalue()
    def getvalue(self):return self.data


def test_systematic_sections_folder_count_and_repeats(tmp_path):
    upload=Upload([
        ('Example vendor/Section A/01 - Alpha Ltd/GST/a.txt',b'a'),
        ('Example vendor/Section A/01 - Alpha Ltd/GST/b.txt',b'a'),
        ('Example vendor/Section B/02 - Beta Ltd/PAN Card.txt',b'b'),
        ('Example vendor/Section B/03 - Empty Ltd/',b'')])
    store=Store(tmp_path)
    result=import_documents(store,[upload])
    assert result.detected_companies==3 and result.processed_files==3
    assert result.saved_files==2 and result.duplicate_files==1
    assert result.total_companies==3 and not result.issues
    again=import_documents(Store(tmp_path),[upload])
    assert again.saved_files==0 and again.duplicate_files==3
    assert dashboard_counts(store.vendors(),store.documents())['companies']==3
    assert len(store.history())==2


def test_reset_history_clears_all_data_including_zip_archive(tmp_path):
    store=Store(tmp_path)
    upload=Upload([('Batch/Alpha Ltd/GST.pdf',b'a')])
    upload.name='Example Vendor Batch.zip'
    import_documents(store,[upload])
    store.upsert_vendors(pd.DataFrame({'company_name':['Empty Vendor']}))
    backup=store.backup_bytes()
    with store.connection() as db:
        db.execute("INSERT INTO vdd_backups VALUES(?,?,?)", ('manual','now',backup))
    with pytest.raises(ValueError):
        store.reset_data('reset')
    assert len(store.documents())==1 and len(store.upload_archives())==1
    store.reset_data('RESET HISTORY')
    reopened=Store(tmp_path)
    assert reopened.vendors().empty and reopened.documents().empty
    assert reopened.history()==[] and reopened.backup_list()==[]
    assert reopened.upload_archives().empty
    assert list(reopened.uploads_dir.glob("*.zip")) == []


def test_reset_history_succeeds_even_if_classified_file_is_missing(tmp_path):
    store=Store(tmp_path)
    store.save_document('Alpha/GST.pdf',b'test',classify('Alpha/GST.pdf'))
    store.safe_path(store.documents().iloc[0].stored_path).unlink()
    store.reset_data('RESET HISTORY')
    assert store.documents().empty and store.vendors().empty
    assert store.history()==[] and store.backup_list()==[]


def test_company_named_export_and_isolation(tmp_path):
    store=Store(tmp_path)
    for name in ['Alpha & Sons','Beta']:
        store.save_document(f'{name}/GST.pdf',name.encode(),classify(f'{name}/GST.pdf'))
    assert export_filename('Alpha & Sons','Documents','zip')=='Alpha_&_Sons_Documents.zip'
    payload=store.company_zip('Alpha & Sons',store.documents())
    with zipfile.ZipFile(BytesIO(payload)) as z:
        assert all(p.startswith('Alpha_&_Sons/') for p in z.namelist())
        for p in z.namelist():
            if p.endswith('.pdf'):assert z.read(p)==b'Alpha & Sons'
            if p.endswith('.xlsx'):
                book=load_workbook(BytesIO(z.read(p)))
                assert 'Alpha & Sons' in book['Document Checklist']['A1'].value
                assert book['Document Checklist']['A5'].value=='Alpha & Sons'
                assert book['Document Checklist'].max_row==5


def test_supporting_documents_visible_in_company_excel(tmp_path):
    store=Store(tmp_path)
    decision=classify('Alpha/Certificate of Incorporation.pdf')
    assert not decision.needs_review and decision.document_types==()
    store.save_document('Alpha/Certificate of Incorporation.pdf',b'incorporation',decision)
    frame=build_checklist(store.vendors(),store.documents())
    assert frame.iloc[0].Available==0 and frame.iloc[0]['Needs review']==0
    book=load_workbook(BytesIO(workbook_bytes(frame,store.documents(),True)))
    assert book['Alpha']['A13'].value=='Incorporation'
    assert book['Alpha']['C13'].value=='Certificate of Incorporation.pdf'


def test_malicious_backup_is_rejected_before_write(tmp_path):
    store=Store(tmp_path);out=BytesIO()
    manifest={'format':'vendor-dashboard-backup-v3','vendors':[{'company_name':'Bad'}], 'documents':[{'file_hash':'../unsafe','backup_member':'../unsafe'}]}
    with zipfile.ZipFile(out,'w') as z:z.writestr('manifest.json',json.dumps(manifest))
    with pytest.raises(ValueError):store.restore_backup(out.getvalue())
    assert store.vendors().empty

def test_vendor_head_count_and_uploaded_zip_archive(tmp_path):
    first=Upload([
        ('Team vendor/01 - Alpha Ltd/GST.pdf',b'a'),
        ('Team vendor/02 - Beta Ltd/PAN Card.pdf',b'b')])
    second=Upload([
        ('New handover/01 - Beta Ltd/GST.pdf',b'c'),
        ('New handover/02 - Gamma LLP/ISO 13485.pdf',b'd')])
    first.name='first_vendor_batch.zip'
    second.name='second_vendor_batch.zip'
    store=Store(tmp_path)
    r1=import_documents(store,[first])
    assert r1.detected_companies==2 and r1.total_companies==2
    archives=store.upload_archives()
    assert len(archives)==1 and int(archives.iloc[0].detected_companies)==2
    assert store.read_upload_archive(archives.iloc[0].id)==first.getvalue()

    r2=import_documents(store,[second])
    assert r2.detected_companies==2 and r2.total_companies==3
    assert dashboard_counts(store.vendors(),store.documents())['companies']==3
    archives=store.upload_archives()
    assert len(archives)==2
    second_row=archives[archives.filename=='second_vendor_batch.zip'].iloc[0]
    assert int(second_row.detected_companies)==2
    assert set(second_row.company_names)=={'Beta Ltd','Gamma LLP'}

    r3=import_documents(store,[second])
    assert r3.total_companies==3
    archives=store.upload_archives()
    assert len(archives)==2
    assert int(archives[archives.filename=='second_vendor_batch.zip'].iloc[0].upload_count)==2


def test_uploaded_zip_archive_is_deleted_by_reset(tmp_path):
    upload=Upload([('Batch/Alpha Ltd/GST.pdf',b'a')])
    upload.name='saved_batch.zip'
    store=Store(tmp_path)
    import_documents(store,[upload])
    before=store.upload_archives().iloc[0]
    stored_path=store._upload_path(before.stored_path)
    assert stored_path and stored_path.is_file()
    store.reset_data('RESET HISTORY')
    assert store.upload_archives().empty
    assert not stored_path.exists()
